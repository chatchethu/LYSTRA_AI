from typing import List, Dict, Optional, Any
from pydantic import BaseModel
from pathlib import Path

class ParsedDocument(BaseModel):
    filename: str
    content_type: str
    text: str
    sections: List[Dict[str, Any]]
    tables: List[Dict[str, Any]]
    metadata: Dict[str, Any]
    page_count: Optional[int]

class DocumentParser:
    """
    Parses various document types into standard representation.
    """
    
    async def parse(self, file_path: str, content_type: str) -> ParsedDocument:
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf" or content_type == "application/pdf":
            return await self.parse_pdf(file_path)
        elif ext == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return await self.parse_docx(file_path)
        elif ext == ".csv" or content_type == "text/csv":
            return await self.parse_csv(file_path)
        elif ext == ".xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return await self.parse_xlsx(file_path)
        elif ext == ".txt" or content_type == "text/plain":
            return await self.parse_txt(file_path)
        elif ext in [".html", ".htm"] or content_type == "text/html":
            with open(file_path, "r", encoding="utf-8") as f:
                return await self.parse_html(f.read())
        else:
            raise ValueError(f"Unsupported document type: {ext} / {content_type}")

    async def parse_pdf(self, file_path: str) -> ParsedDocument:
        try:
            import PyPDF2
            text = ""
            page_count = 0
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                page_count = len(reader.pages)
                for page_num in range(page_count):
                    page = reader.pages[page_num]
                    text += page.extract_text() + "\n"
                    
            return ParsedDocument(
                filename=Path(file_path).name,
                content_type="application/pdf",
                text=text,
                sections=[],
                tables=[],
                metadata={"source": file_path},
                page_count=page_count
            )
        except ImportError:
            raise ImportError("PyPDF2 is required for PDF parsing")

    async def parse_docx(self, file_path: str) -> ParsedDocument:
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return ParsedDocument(
                filename=Path(file_path).name,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                text=text,
                sections=[],
                tables=[],
                metadata={"source": file_path},
                page_count=None
            )
        except ImportError:
            raise ImportError("python-docx is required for docx parsing")

    async def parse_xlsx(self, file_path: str) -> ParsedDocument:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path)
            text = ""
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text += f"Sheet: {sheet}\n"
                for row in ws.iter_rows(values_only=True):
                    text += "\t".join([str(c) if c is not None else "" for c in row]) + "\n"
                    
            return ParsedDocument(
                filename=Path(file_path).name,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                text=text,
                sections=[],
                tables=[],
                metadata={"source": file_path},
                page_count=None
            )
        except ImportError:
            raise ImportError("openpyxl is required for xlsx parsing")

    async def parse_csv(self, file_path: str) -> ParsedDocument:
        import csv
        text = ""
        with open(file_path, newline='', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                text += ",".join(row) + "\n"
                
        return ParsedDocument(
            filename=Path(file_path).name,
            content_type="text/csv",
            text=text,
            sections=[],
            tables=[],
            metadata={"source": file_path},
            page_count=None
        )

    async def parse_txt(self, file_path: str) -> ParsedDocument:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        return ParsedDocument(
            filename=Path(file_path).name,
            content_type="text/plain",
            text=text,
            sections=[],
            tables=[],
            metadata={"source": file_path},
            page_count=None
        )

    async def parse_html(self, content: str) -> ParsedDocument:
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(content, "html.parser")
            text = soup.get_text(separator="\n", strip=True)
            return ParsedDocument(
                filename="html_content.html",
                content_type="text/html",
                text=text,
                sections=[],
                tables=[],
                metadata={},
                page_count=None
            )
        except ImportError:
            raise ImportError("beautifulsoup4 is required for HTML parsing")
