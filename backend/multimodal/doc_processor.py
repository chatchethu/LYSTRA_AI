"""
Document Processor — LYSTRA Multimodal V1

Handles DOCX files using python-docx:
  - Extracts headings, paragraphs, tables
  - Sends embedded images to vision model
"""
from __future__ import annotations
import io
from typing import TYPE_CHECKING

try:
    import docx as _docx
    _DOCX_OK = True
except ImportError:
    _DOCX_OK = False

from backend.multimodal.file_context import FileContext, FileChunk, FileType, FileStatus

if TYPE_CHECKING:
    from backend.llm.gateway import LLMGateway


class DocProcessor:
    """Extracts structured content from DOCX files."""

    def __init__(self, llm: "LLMGateway | None" = None, vision_model: str = "qwen2.5vl:3b"):
        self.llm = llm
        self.vision_model = vision_model

    async def process(self, file_id: str, filename: str, file_bytes: bytes) -> FileContext:
        if not _DOCX_OK:
            return FileContext(
                file_id=file_id, filename=filename, file_type=FileType.DOCUMENT,
                status=FileStatus.FAILED, error="python-docx not installed"
            )

        try:
            doc = _docx.Document(io.BytesIO(file_bytes))
            chunks: list[FileChunk] = []
            current_heading = None
            para_num = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                style_name = para.style.name.lower() if para.style else ""
                is_heading = "heading" in style_name

                if is_heading:
                    current_heading = text
                    chunks.append(FileChunk(
                        chunk_type="heading",
                        content=text,
                        section=text
                    ))
                else:
                    para_num += 1
                    chunks.append(FileChunk(
                        chunk_type="text",
                        content=text,
                        section=current_heading
                    ))

            # Extract tables
            for table_idx, table in enumerate(doc.tables, start=1):
                rows = []
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    table_text = "\n".join(rows)
                    chunks.append(FileChunk(
                        chunk_type="table",
                        content=f"Table {table_idx}:\n{table_text}",
                        section=current_heading
                    ))

            # Extract inline images → vision model
            if self.llm:
                image_rels = [
                    rel for rel in doc.part.rels.values()
                    if "image" in rel.reltype
                ]
                for img_idx, rel in enumerate(image_rels[:5], start=1):  # cap at 5 images
                    try:
                        img_bytes = rel.target_part.blob
                        desc = await self._describe_image(img_bytes)
                        if desc:
                            chunks.append(FileChunk(
                                chunk_type="visual_analysis",
                                content=f"[Image {img_idx}]: {desc}",
                                section=current_heading
                            ))
                    except Exception:
                        pass

            if not chunks:
                return FileContext(
                    file_id=file_id, filename=filename, file_type=FileType.DOCUMENT,
                    status=FileStatus.FAILED, error="No readable content found in document"
                )

            return FileContext(
                file_id=file_id, filename=filename, file_type=FileType.DOCUMENT,
                status=FileStatus.READY, chunks=chunks
            )

        except Exception as e:
            return FileContext(
                file_id=file_id, filename=filename, file_type=FileType.DOCUMENT,
                status=FileStatus.FAILED, error=str(e)
            )

    async def _describe_image(self, image_bytes: bytes) -> str:
        """Send embedded image to vision model for description."""
        import base64
        b64 = base64.b64encode(image_bytes).decode()
        messages = [{
            "role": "user",
            "content": "Describe what is shown in this image concisely.",
            "images": [b64]
        }]
        try:
            return await self.llm.chat(messages=messages, model=self.vision_model)
        except Exception:
            return ""
